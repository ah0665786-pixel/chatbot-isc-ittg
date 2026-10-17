from __future__ import annotations

from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


def load_sections(path: Path) -> list[tuple[str, str]]:
    """Devuelve [(sección, texto)]: una por página en PDF, una sola en el resto."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [(f"pág. {i}", (p.extract_text() or "")) for i, p in enumerate(reader.pages, 1)]
    if ext == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:  # los reglamentos suelen traer tablas
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return [("", "\n\n".join(parts))]
    if ext in {".txt", ".md"}:
        return [("", path.read_text(encoding="utf-8", errors="replace"))]
    raise ValueError(f"formato no soportado: {ext}")
