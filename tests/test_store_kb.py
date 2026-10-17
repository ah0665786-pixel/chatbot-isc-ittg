import numpy as np
import pytest

from app.adapters.embedders import HashingEmbedder
from app.adapters.vectorstore_faiss import FaissStore, IncompatibleIndex
from app.domain import Chunk
from app.rag.knowledge_base import KnowledgeBase


def _chunks(doc, n):
    return [Chunk(doc, doc + ".txt", i, f"texto {doc} {i}") for i in range(n)]


def test_store_search_remove_persist(tmp_path):
    e = HashingEmbedder()
    s = FaissStore(e.dim, e.name)
    a, b = _chunks("a", 2), _chunks("b", 2)
    s.add(a, e.embed_documents([c.text for c in a]))
    s.add(b, e.embed_documents(["baja temporal semestre", "residencia profesional horas"]))
    hits = s.search(e.embed_query("residencia profesional"), 2)
    assert hits[0].chunk.doc_id == "b" and 0 < hits[0].score <= 1.0001
    assert s.remove_document("a") == 2 and len(s) == 2
    s.save(tmp_path)
    s2 = FaissStore.load(tmp_path, e.dim, e.name)
    assert len(s2) == 2
    assert s2.search(e.embed_query("residencia profesional"), 1)[0].chunk.doc_id == "b"


def test_incompatible_embedder_refuses_to_load(tmp_path):
    e = HashingEmbedder()
    FaissStore(e.dim, e.name).save(tmp_path)
    with pytest.raises(IncompatibleIndex):
        FaissStore.load(tmp_path, 384, "otro-modelo")


def test_ingest_added_unchanged_replaced_and_persisted(cfg, embedder, tmp_path):
    docs = tmp_path / "d"
    docs.mkdir()
    f = docs / "Manual ISC.txt"
    f.write_text("El servicio social requiere 480 horas. " * 30, encoding="utf-8")
    kb = KnowledgeBase.open(cfg, embedder)
    assert kb.ingest_file(f, "v1").status == "added"
    assert kb.ingest_file(f, "v1").status == "unchanged"
    n = len(kb.store)
    f.write_text("El servicio social requiere 500 horas. " * 30, encoding="utf-8")
    assert kb.ingest_file(f, "v2").status == "replaced"
    assert len(kb.store) == n and kb.docs["manual-isc"].version == "v2"
    kb.save()
    kb2 = KnowledgeBase.open(cfg, embedder)
    assert len(kb2.store) == n and "manual-isc" in kb2.docs


def test_ingest_empty_document_reports_error(cfg, embedder, tmp_path):
    f = tmp_path / "vacio.txt"
    f.write_text("   ", encoding="utf-8")
    r = KnowledgeBase.open(cfg, embedder).ingest_file(f)
    assert r.status == "error" and "sin texto" in r.message


def test_ingest_docx(cfg, embedder, tmp_path):
    from docx import Document

    d = Document()
    d.add_paragraph("La titulación integral requiere liberación del servicio social.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Requisito", "Constancia de inglés"
    f = tmp_path / "Titulacion.docx"
    d.save(f)
    kb = KnowledgeBase.open(cfg, embedder)
    assert kb.ingest_file(f).status == "added"
    txt = " ".join(c.text for c in kb.store._chunks)
    assert "Constancia de inglés" in txt
